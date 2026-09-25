"""Score models on a reviewed sheet from draft_labels.py. Every non-empty question cell in the `review` sheet is taken as the
human label (the reviewer corrected the teacher's draft). Each model endpoint is asked the same questions and scored per
question: accuracy for choice/noul, exact-level accuracy + MAE for score. Also writes the eval set in eval_thai.py format
(labels as strings / bools / level indexes) so future checkpoints can be scored offline on the GPU box.

    python thai/real_eval/score_real.py --sheet tickets_review.xlsx --out thai/real_eval/tickets_eval.jsonl \
        --model student=http://172.18.72.145:8011 --model teacher=http://172.18.72.145:8010
"""
import argparse
import json
import os
import sys
import urllib.request
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from cc_questions import QUESTION_SETS, get_questions  # noqa: E402

QUESTIONS, ORDER = get_questions("ecom")


def ask(url, text, qs):
    body = json.dumps({"state": text, "questions": qs}, ensure_ascii=False).encode()
    req = urllib.request.Request(url + "/v1/systemone", body, {"content-type": "application/json"})
    for _ in range(3):
        try:
            return json.load(urllib.request.urlopen(req, timeout=120))["answers"]
        except Exception:  # noqa: BLE001
            pass
    return None


def parse_label(q, cell):
    """sheet value -> label in eval_thai.py form (choice: key, noul: bool, score: level index) or None if not labelled"""
    if cell is None or str(cell).strip() == "":
        return None
    v = str(cell).strip()
    if q["type"] == "choice":
        return v if v in q["criteria"] else None
    if q["type"] == "noul":
        return v.lower() in ("yes", "true", "1", "ใช่")
    return q["criteria"].index(v) if v in q["criteria"] else (int(v) if v.isdigit() and int(v) < len(q["criteria"]) else None)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--out", help="write the labelled records here (eval_thai.py format)")
    ap.add_argument("--model", action="append", default=[], help="name=url, repeatable")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--questions", default="ecom", choices=sorted(QUESTION_SETS), help="question set the sheet was drafted with")
    args = ap.parse_args()
    global QUESTIONS, ORDER
    QUESTIONS, ORDER = get_questions(args.questions)

    ws = load_workbook(args.sheet, read_only=True)["review"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    col = {h: i for i, h in enumerate(header)}
    records = []
    for r in rows[1:]:
        if not r or not r[col["text"]]:
            continue
        labels = {}
        for q in ORDER:
            lab = parse_label(QUESTIONS[q], r[col[q]])
            if lab is not None:
                labels[q] = lab
        if labels:
            records.append({"id": f"real-{len(records)}", "source": "real", "state": str(r[col["text"]]).strip(),
                            "questions": {q: QUESTIONS[q] for q in labels}, "labels": labels})
    n_lab = sum(len(r["labels"]) for r in records)
    print(f"{len(records)} labelled texts, {n_lab} labelled decisions", flush=True)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            for r in records:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")

    for spec in args.model:
        name, url = spec.split("=", 1)
        with ThreadPoolExecutor(args.workers) as ex:
            answers = list(ex.map(lambda r: ask(url, r["state"], r["questions"]), records))
        st = defaultdict(lambda: {"n": 0, "ok": 0, "mae": 0.0})
        for r, a in zip(records, answers):
            if a is None:
                continue
            for q, lab in r["labels"].items():
                s, t = st[q], QUESTIONS[q]["type"]
                s["n"] += 1
                if t == "choice":
                    s["ok"] += a[q]["choice"] == lab
                elif t == "noul":
                    s["ok"] += (a[q]["noul"] > 0.5) == lab
                else:
                    s["ok"] += round(a[q]["score"]) == lab
                    s["mae"] += abs(a[q]["score"] - lab)
        tot_n = sum(s["n"] for s in st.values())
        tot_ok = sum(s["ok"] for s in st.values())
        print(f"\n== {name} ({url}): overall {tot_ok}/{tot_n} = {tot_ok / max(1, tot_n):.3f}")
        print(f"{'question':14s} {'n':>4s} {'acc':>6s} {'mae':>5s}")
        for q in ORDER:
            if q in st:
                s = st[q]
                mae = f"{s['mae'] / s['n']:5.2f}" if QUESTIONS[q]["type"] == "score" else "     "
                print(f"{q:14s} {s['n']:4d} {s['ok'] / s['n']:6.3f} {mae}")


if __name__ == "__main__":
    main()
