"""Turn real call-center texts into a review sheet: the teacher (OpenThai-SystemOne) drafts every label, the student
(cascade endpoint, run 5) answers alongside, and a reviewer corrects the draft columns in Excel.

Input: a .txt (one text per line), .csv (column `text`) or .jsonl (`text` field). Output: an .xlsx with one row per text,
one column per question (dropdowns for choice/score/noul), prefilled with the teacher's answer, plus hidden helper columns
with the student's answer and both confidences so disagreements can be reviewed first (sorted to the top).

    python thai/real_eval/draft_labels.py --inp tickets.txt --out tickets_review.xlsx
"""
import argparse
import csv
import io
import json
import os
import sys
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from cc_questions import QUESTION_SETS, get_questions  # noqa: E402  ("ecom": the 9-question call-center set with `other`; "debt": debt collection)

QUESTIONS, ORDER = get_questions("ecom")


def read_texts(path):
    if path.endswith(".jsonl"):
        texts = [json.loads(l)["text"] for l in open(path, encoding="utf-8") if l.strip()]
    else:
        raw = open(path, encoding="utf-8-sig").read()
        first = raw.splitlines()[0].strip().lower() if raw.strip() else ""
        if path.endswith(".csv") or first == "text" or first.startswith("text,"):  # a CSV with a `text` column, whatever the extension
            texts = [r["text"] for r in csv.DictReader(io.StringIO(raw)) if r.get("text")]
        else:
            texts = [l.strip() for l in raw.splitlines() if l.strip()]
    seen, out = set(), []
    for t in texts:  # drop exact duplicates, keep order
        t = t.strip()
        if t and t not in seen:
            seen.add(t)
            out.append(t)
    return out


def ask(url, text, tries=5):
    # single option order (order_invariant off): the teacher would otherwise average 28 permutations for the intent question
    # (~2x slower, timeouts under load), and the student's labels were produced in single order too
    body = json.dumps({"state": text, "questions": QUESTIONS, "order_invariant": False}, ensure_ascii=False).encode()
    req = urllib.request.Request(url + "/v1/systemone", body, {"content-type": "application/json"})
    for attempt in range(tries):
        try:
            return json.load(urllib.request.urlopen(req, timeout=180))["answers"]
        except Exception:  # noqa: BLE001
            time.sleep(2 * (attempt + 1))
    return None


def value(q, a):
    """answer -> the label the reviewer sees, and the model's confidence in it"""
    if q["type"] == "choice":
        return a["choice"], max(a["probabilities"].values())
    if q["type"] == "noul":
        return ("yes" if a["noul"] > 0.5 else "no"), max(a["noul"], 1 - a["noul"])
    idx = round(a["score"])
    return q["criteria"][idx], a["probabilities"][str(idx)]


def options(q):
    if q["type"] == "choice":
        return list(q["criteria"])
    if q["type"] == "noul":
        return ["yes", "no"]
    return list(q["criteria"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inp", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--teacher", default="http://172.18.72.145:8010")
    ap.add_argument("--student", default="http://172.18.72.145:8011")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--questions", default="ecom", choices=sorted(QUESTION_SETS), help="question set: ecom (call-center, default) or debt (debt collection)")
    args = ap.parse_args()
    global QUESTIONS, ORDER
    QUESTIONS, ORDER = get_questions(args.questions)
    texts = read_texts(args.inp)
    print(f"{len(texts)} texts", flush=True)
    with ThreadPoolExecutor(args.workers) as ex:
        t_ans = list(ex.map(lambda t: ask(args.teacher, t), texts))
        s_ans = list(ex.map(lambda t: ask(args.student, t), texts))

    rows = []
    for text, ta, sa in zip(texts, t_ans, s_ans):
        row = {"text": text, "disagreements": 0}
        for qid in ORDER:
            q = QUESTIONS[qid]
            tv, tc = value(q, ta[qid]) if ta else ("", 0)
            sv, sc = value(q, sa[qid]) if sa else ("", 0)
            row[qid] = tv
            row[qid + "__student"] = sv
            row[qid + "__teacher_p"] = round(tc, 2)
            row[qid + "__student_p"] = round(sc, 2)
            row["disagreements"] += int(tv != sv)
        rows.append(row)
    rows.sort(key=lambda r: -r["disagreements"])  # review the contested ones first

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    header = ["#", "text"] + ORDER + ["disagreements", "reviewer_note"] + [f"{q}__student" for q in ORDER] + [f"{q}__teacher_p" for q in ORDER] + [f"{q}__student_p" for q in ORDER]
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for i, r in enumerate(rows, start=2):
        ws.append([i - 1, r["text"]] + [r[q] for q in ORDER] + [r["disagreements"], ""] + [r[q + "__student"] for q in ORDER]
                  + [r[q + "__teacher_p"] for q in ORDER] + [r[q + "__student_p"] for q in ORDER])
        for j, q in enumerate(ORDER):
            if r[q] != r[q + "__student"]:
                ws.cell(row=i, column=3 + j).fill = fill  # highlight where teacher and student disagree
    ws.column_dimensions["B"].width = 70
    for j, q in enumerate(ORDER):
        col = get_column_letter(3 + j)
        ws.column_dimensions[col].width = max(14, len(q) + 2)
        dv = DataValidation(type="list", formula1='"' + ",".join(options(QUESTIONS[q])) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(rows) + 1}")
    for c in range(1, len(header) + 1):
        ws.cell(row=1, column=c).alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "C2"
    for c in range(3 + len(ORDER) + 2, len(header) + 1):  # helper columns: hide
        ws.column_dimensions[get_column_letter(c)].hidden = True
    legend = wb.create_sheet("questions")
    legend.append(["question", "type", "instructions", "options"])
    for q in ORDER:
        crit = QUESTIONS[q].get("criteria")
        if isinstance(crit, dict):
            desc = " | ".join(f"{k}: {v}" for k, v in crit.items())
        elif crit:
            desc = " | ".join(f"{i}: {v}" for i, v in enumerate(crit))
        else:
            desc = "yes | no"
        legend.append([q, QUESTIONS[q]["type"], QUESTIONS[q]["instructions"], desc])
    legend.column_dimensions["C"].width = 50
    legend.column_dimensions["D"].width = 120
    wb.save(args.out)
    n_dis = sum(1 for r in rows if r["disagreements"])
    print(f"wrote {args.out}: {len(rows)} rows, {n_dis} with at least one teacher/student disagreement (highlighted, sorted first)")


if __name__ == "__main__":
    main()
